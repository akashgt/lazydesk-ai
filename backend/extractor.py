import os
from pathlib import Path
from typing import Optional, List, Dict, Any
from pydantic import BaseModel
import fitz  # PyMuPDF
import docx
from backend.config import settings

class ExtractedDocument(BaseModel):
    file_path: str
    file_name: str
    file_size_bytes: int
    extension: str
    word_count: int
    page_count: Optional[int] = None
    extracted_text: str
    chunks: List[str]
    has_text: bool
    error: Optional[str] = None

def chunk_text(text: str, max_chunk_chars: int = settings.max_chunk_chars) -> List[str]:
    """Splits text into readable chunks without cutting words across boundaries."""
    if not text:
        return []
    
    words = text.split()
    chunks: List[str] = []
    current_chunk: List[str] = []
    current_length = 0

    for word in words:
        word_len = len(word) + 1
        if current_length + word_len > max_chunk_chars:
            if current_chunk:
                chunks.append(" ".join(current_chunk))
                current_chunk = [word]
                current_length = word_len
            else:
                chunks.append(word)
                current_chunk = []
                current_length = 0
        else:
            current_chunk.append(word)
            current_length += word_len

    if current_chunk:
        chunks.append(" ".join(current_chunk))

    return chunks

def extract_text_from_pdf(file_path: Path) -> tuple[str, int]:
    """Extracts text and page count from a PDF using PyMuPDF."""
    doc = fitz.open(file_path)
    page_count = len(doc)
    text_parts = []
    for page in doc:
        text_parts.append(page.get_text("text"))
    doc.close()
    return "\n\n".join(text_parts).strip(), page_count

def extract_text_from_docx(file_path: Path) -> tuple[str, int]:
    """Extracts text from a DOCX file using python-docx."""
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs).strip(), 1

def extract_text_from_plain(file_path: Path) -> tuple[str, int]:
    """Extracts text from plaintext, markdown, csv, or code files."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
    except Exception:
        with open(file_path, "r", encoding="latin-1", errors="replace") as f:
            content = f.read()
    return content.strip(), 1

def extract_document(file_path: Path) -> ExtractedDocument:
    """Safely extracts text, statistics, and chunks from a supported document."""
    resolved_path = file_path.resolve()
    file_name = resolved_path.name
    ext = resolved_path.suffix.lower()

    if not resolved_path.exists() or not resolved_path.is_file():
        return ExtractedDocument(
            file_path=str(resolved_path),
            file_name=file_name,
            file_size_bytes=0,
            extension=ext,
            word_count=0,
            extracted_text="",
            chunks=[],
            has_text=False,
            error="File does not exist or is not a regular file.",
        )

    size = resolved_path.stat().st_size
    if size > settings.max_extract_size_bytes:
        return ExtractedDocument(
            file_path=str(resolved_path),
            file_name=file_name,
            file_size_bytes=size,
            extension=ext,
            word_count=0,
            extracted_text="",
            chunks=[],
            has_text=False,
            error=f"File exceeds maximum allowed extraction size ({size / (1024 * 1024):.1f} MB > {settings.max_extract_size_bytes / (1024 * 1024)} MB).",
        )

    text = ""
    pages = 1
    err = None

    try:
        if ext == ".pdf":
            text, pages = extract_text_from_pdf(resolved_path)
        elif ext == ".docx":
            text, pages = extract_text_from_docx(resolved_path)
        elif ext in {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".py", ".js", ".ts", ".html", ".css", ".rtf"}:
            text, pages = extract_text_from_plain(resolved_path)
        else:
            err = f"Unsupported document format: {ext}"
    except Exception as e:
        err = f"Failed to extract document contents: {str(e)}"

    words = len(text.split()) if text else 0
    chunks = chunk_text(text) if text else []

    return ExtractedDocument(
        file_path=str(resolved_path),
        file_name=file_name,
        file_size_bytes=size,
        extension=ext,
        word_count=words,
        page_count=pages,
        extracted_text=text,
        chunks=chunks,
        has_text=bool(text.strip()),
        error=err,
    )
